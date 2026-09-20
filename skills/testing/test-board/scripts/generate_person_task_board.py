from __future__ import annotations

"""从统一任务总表生成按人员维度的任务看板。

本脚本读取测试任务总表 Excel，按负责人拆分任务条目，生成静态 HTML 任务看板和可下载的 Excel 看板。
"""

import argparse
import html
import json
import re
import shutil
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OWNER_SPLIT_RE = re.compile(r"[、，,/\s]+")
DAY_COLUMN_WIDTH = 132
PERSON_COLUMN_WIDTH = 150
TEMPLATE_PATH = Path(__file__).resolve().with_name("templates") / "person_task_board.html"
_SCRIPT_DIR = Path(__file__).resolve().parent
_WORKSPACE_ROOT = Path.cwd()
DEFAULT_OUTPUT_DIR = _WORKSPACE_ROOT
INDEX_OUTPUT_NAME = "index.html"
PENDING_SCHEDULE_STATUS = "待排期"

TASK_BAR_THEME = {
    "bar": "#3370FF",
    "bar_deep": "#3370FF",
    "today": "#FFFFFF",
    "weekend": "#F3F4F6",
    "idle": "#FCA5A5",
    "grid": "#E6EAF0",
    "grid_strong": "#D8DEE9",
    "case_tag": "rgba(186, 206, 253, 0.78)",
    "execution_tag": "rgba(51, 112, 255, 0.26)",
    "expected_tag": "rgba(51, 112, 255, 0.18)",
}
HOLIDAY_FILL = "#F3F4F6"
HOLIDAY_LABELS = {
    date(2026, 6, 19): "端午",
    date(2026, 6, 20): "端午",
    date(2026, 6, 21): "端午",
}
REQUIRED_HEADERS = [
    "任务编号",
    "任务标题",
    "负责人",
    "任务来源",
    "业务域",
    "测试类型",
    "优先级",
    "所属版本",
    "当前状态",
    "预计提测时间",
    "用例开始时间",
    "用例结束时间",
    "计划开始时间",
    "计划完成时间",
    "备注",
    "父记录",
]


@dataclass
class Task:
    task_id: str
    title: str
    owners: list[str]
    source: str
    domain: str
    test_type: str
    priority: str
    version: str
    status: str
    expected_test_date: date | None
    case_start: date | None
    case_end: date | None
    plan_start: date | None
    plan_end: date | None
    remark: str
    parent_record: str

@dataclass
class PersonTask:
    owner: str
    task: Task
    start: date
    end: date
    segment_type: str
    segment_label: str

    @property
    def span(self) -> int:
        return (self.end - self.start).days + 1


def parse_args() -> argparse.Namespace:
    default_input = _WORKSPACE_ROOT / "测试管理甘特图_测试部门统一任务总表.xlsx"
    default_output = DEFAULT_OUTPUT_DIR / "人员任务看板.html"
    default_excel_output = DEFAULT_OUTPUT_DIR / "人员任务看板.xlsx"
    parser = argparse.ArgumentParser(description="从统一任务总表生成按人员维度的静态任务网页。")
    parser.add_argument("--input", type=Path, default=default_input, help="输入 Excel 路径")
    parser.add_argument("--output", type=Path, default=default_output, help="输出 HTML 路径")
    parser.add_argument("--excel-output", type=Path, default=default_excel_output, help="输出人员任务看板 Excel 路径")
    parser.add_argument("--start-date", type=str, default="", help="人员任务看板开始日期，格式 YYYY-MM-DD")
    parser.add_argument("--end-date", type=str, default="", help="人员任务看板截止日期，格式 YYYY-MM-DD")
    parser.add_argument("--excel-start", type=str, default="", help="兼容旧参数：等同 --start-date")
    parser.add_argument("--excel-end", type=str, default="", help="兼容旧参数：等同 --end-date")
    return parser.parse_args()


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def to_date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_date_arg(value: str, arg_name: str) -> date | None:
    if not value:
        return None
    parsed = to_date(value)
    if parsed is None:
        raise ValueError(f"{arg_name} 日期格式无效: {value}，应为 YYYY-MM-DD")
    return parsed


def split_owners(value: str) -> list[str]:
    owners = []
    for raw in OWNER_SPLIT_RE.split(value):
        name = raw.strip()
        if name and name not in owners:
            owners.append(name)
    return owners


def load_tasks(input_path: Path) -> list[Task]:
    # Some workbooks report an incorrect used range in openpyxl read-only mode.
    # Normal mode reliably reads the actual rows/columns for this task sheet.
    workbook = load_workbook(input_path, read_only=False, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [normalize_text(cell) for cell in next(rows)]
    header_index = {name: idx for idx, name in enumerate(headers)}

    missing = [name for name in REQUIRED_HEADERS if name not in header_index]
    if missing:
        raise ValueError(f"Excel 缺少必要列: {', '.join(missing)}")

    tasks: list[Task] = []
    for row in rows:
        title = normalize_text(row[header_index["任务标题"]])
        if not title:
            continue
        owners = split_owners(normalize_text(row[header_index["负责人"]]))

        tasks.append(
            Task(
                task_id=normalize_text(row[header_index["任务编号"]]),
                title=title,
                owners=owners,
                source=normalize_text(row[header_index["任务来源"]]),
                domain=normalize_text(row[header_index["业务域"]]),
                test_type=normalize_text(row[header_index["测试类型"]]),
                priority=normalize_text(row[header_index["优先级"]]),
                version=normalize_text(row[header_index["所属版本"]]),
                status=normalize_text(row[header_index["当前状态"]]),
                expected_test_date=to_date(row[header_index["预计提测时间"]]),
                case_start=to_date(row[header_index["用例开始时间"]]),
                case_end=to_date(row[header_index["用例结束时间"]]),
                plan_start=to_date(row[header_index["计划开始时间"]]),
                plan_end=to_date(row[header_index["计划完成时间"]]),
                remark=normalize_text(row[header_index["备注"]]),
                parent_record=normalize_text(row[header_index["父记录"]]),
            )
        )
    return tasks


def daterange(start: date, end: date) -> Iterable[date]:
    current = start
    while current <= end:
        yield current
        current += timedelta(days=1)


def is_holiday(day: date) -> bool:
    return day in HOLIDAY_LABELS


def get_day_label(day: date) -> str:
    weekday = "一二三四五六日"[day.weekday()]
    holiday_label = HOLIDAY_LABELS.get(day)
    if holiday_label:
        return f"{day.day}\n周{weekday}\n{holiday_label}"
    return f"{day.day}\n周{weekday}"


def priority_rank(value: str) -> int:
    mapping = {"P0": 0, "P1": 1, "高": 1, "紧急": 1, "P2": 2, "中": 2, "P3": 3, "低": 3}
    return mapping.get(value.upper() if value else "", 9)


def normalize_range(start: date | None, end: date | None) -> tuple[date | None, date | None]:
    resolved_start = start or end
    resolved_end = end or start
    if resolved_start and resolved_end and resolved_end < resolved_start:
        resolved_start, resolved_end = resolved_end, resolved_start
    return resolved_start, resolved_end


def build_task_segments(task: Task) -> list[tuple[date, date, str, str]]:
    segments: list[tuple[date, date, str, str]] = []

    case_start, case_end = normalize_range(task.case_start, task.case_end)
    if case_start and case_end:
        segments.append((case_start, case_end, "case", "测试用例编写"))

    exec_start, exec_end = normalize_range(task.plan_start, task.plan_end)
    if exec_start and exec_end:
        segments.append((exec_start, exec_end, "execution", "测试执行"))
    elif task.expected_test_date:
        expected_start, expected_end = normalize_range(task.expected_test_date, task.expected_test_date)
        if expected_start and expected_end:
            segments.append((expected_start, expected_end, "expected", "预计提测"))

    return segments


def is_pending_schedule(task: Task) -> bool:
    return normalize_text(task.status) == PENDING_SCHEDULE_STATUS


def has_any_schedule_date(task: Task) -> bool:
    return any(
        [
            task.expected_test_date,
            task.case_start,
            task.case_end,
            task.plan_start,
            task.plan_end,
        ]
    )


def filter_display_unscheduled(unscheduled_by_person: dict[str, list[Task]]) -> dict[str, list[Task]]:
    filtered: dict[str, list[Task]] = {}
    for owner, task_list in unscheduled_by_person.items():
        visible_tasks = [task for task in task_list if not is_pending_schedule(task)]
        if visible_tasks:
            filtered[owner] = visible_tasks
    return filtered


def subtract_ranges(
    start: date,
    end: date,
    blockers: list[tuple[date, date]],
) -> list[tuple[date, date]]:
    remaining = [(start, end)]
    for block_start, block_end in sorted(blockers):
        next_remaining: list[tuple[date, date]] = []
        for part_start, part_end in remaining:
            overlap_start = max(part_start, block_start)
            overlap_end = min(part_end, block_end)
            if overlap_start > overlap_end:
                next_remaining.append((part_start, part_end))
                continue
            if part_start < overlap_start:
                next_remaining.append((part_start, overlap_start - timedelta(days=1)))
            if overlap_end < part_end:
                next_remaining.append((overlap_end + timedelta(days=1), part_end))
        remaining = next_remaining
    return remaining


def split_interrupted_task(item: PersonTask, person_tasks: list[PersonTask]) -> list[PersonTask]:
    blockers: list[tuple[date, date]] = []
    for other in person_tasks:
        if other is item:
            continue
        if other.task.task_id == item.task.task_id and other.segment_type == item.segment_type:
            continue
        # A later-starting task is treated as an inserted task and cuts the earlier task.
        if item.start < other.start <= item.end:
            overlap_start = max(item.start, other.start)
            overlap_end = min(item.end, other.end)
            if overlap_start <= overlap_end:
                blockers.append((overlap_start, overlap_end))

    return [
        PersonTask(
            owner=item.owner,
            task=item.task,
            start=part_start,
            end=part_end,
            segment_type=item.segment_type,
            segment_label=item.segment_label,
        )
        for part_start, part_end in subtract_ranges(item.start, item.end, blockers)
    ]


def split_interrupted_lanes(
    lanes: list[list[PersonTask]],
    person_tasks: list[PersonTask],
) -> list[list[PersonTask]]:
    split_lanes: list[list[PersonTask]] = []
    for lane in lanes:
        split_lane: list[PersonTask] = []
        for item in lane:
            split_lane.extend(split_interrupted_task(item, person_tasks))
        split_lanes.append(split_lane)
    return split_lanes


def build_schedule(
    tasks: list[Task],
    window_start: date,
    window_end: date,
) -> tuple[dict[str, list[list[PersonTask]]], dict[str, list[Task]], list[date]]:
    scheduled_by_person: dict[str, list[PersonTask]] = defaultdict(list)
    unscheduled_by_person: dict[str, list[Task]] = defaultdict(list)

    for task in tasks:
        segments = build_task_segments(task)

        for owner in task.owners:
            if segments:
                for start, end, segment_type, segment_label in segments:
                    clipped_start = max(start, window_start)
                    clipped_end = min(end, window_end)
                    if clipped_start <= clipped_end:
                        person_task = PersonTask(
                            owner=owner,
                            task=task,
                            start=clipped_start,
                            end=clipped_end,
                            segment_type=segment_type,
                            segment_label=segment_label,
                        )
                        scheduled_by_person[owner].append(person_task)
            else:
                unscheduled_by_person[owner].append(task)

    lanes_by_person: dict[str, list[list[PersonTask]]] = {}
    for owner, person_tasks in scheduled_by_person.items():
        person_tasks.sort(
            key=lambda item: (
                item.start,
                item.end,
                item.segment_type,
                priority_rank(item.task.priority),
                item.task.task_id,
                item.task.title,
            )
        )
        lanes: list[list[PersonTask]] = []
        lane_ends: list[date] = []
        for person_task in person_tasks:
            placed = False
            for idx, lane_end in enumerate(lane_ends):
                if person_task.start > lane_end:
                    lanes[idx].append(person_task)
                    lane_ends[idx] = person_task.end
                    placed = True
                    break
            if not placed:
                lanes.append([person_task])
                lane_ends.append(person_task.end)
        lanes_by_person[owner] = lanes

    for owner, task_list in unscheduled_by_person.items():
        task_list.sort(key=lambda task: (priority_rank(task.priority), task.task_id, task.title))

    all_dates = list(daterange(window_start, window_end))
    return lanes_by_person, dict(unscheduled_by_person), all_dates


def format_day(day: date) -> str:
    return f"{day.day}"


def render_month_headers(all_dates: list[date]) -> str:
    if not all_dates:
        return ""

    cells: list[str] = []
    current_year = all_dates[0].year
    current_month = all_dates[0].month
    span = 0

    for day in all_dates:
        if day.year == current_year and day.month == current_month:
            span += 1
            continue
        cells.append(
            f'<th class="month-cell sticky-month" colspan="{span}">{current_year}年{current_month}月</th>'
        )
        current_year = day.year
        current_month = day.month
        span = 1

    cells.append(f'<th class="month-cell sticky-month" colspan="{span}">{current_year}年{current_month}月</th>')
    return "".join(cells)


def render_colgroup(all_dates: list[date]) -> str:
    day_cols = "".join(
        [
            f'<col class="day-col" data-date-index="{idx}" data-date="{day.isoformat()}" style="width:{DAY_COLUMN_WIDTH}px">'
            for idx, day in enumerate(all_dates)
        ]
    )
    return f'<colgroup><col class="person-col" style="width:{PERSON_COLUMN_WIDTH}px">{day_cols}</colgroup>'


def middle_ellipsis(text: str, max_chars: int) -> str:
    if max_chars < 5 or len(text) <= max_chars:
        return text
    head_len = (max_chars - 3) // 2
    tail_len = max_chars - 3 - head_len
    return f"{text[:head_len]}...{text[-tail_len:]}"


def build_span_colors(
    start_idx: int,
    span: int,
    all_dates: list[date],
    occupied_indices: set[int] | None = None,
    today_date: date | None = None,
) -> list[str]:
    if span <= 0:
        return ["#ffffff"]

    colors: list[str] = []
    for offset in range(span):
        day_idx = start_idx + offset
        day = all_dates[start_idx + offset]
        should_mark_idle = (
            occupied_indices is not None
            and day_idx not in occupied_indices
            and not is_holiday(day)
            and day.weekday() < 5
        )
        if is_holiday(day):
            color = HOLIDAY_FILL
        elif should_mark_idle:
            color = "var(--idle)"
        else:
            color = "#ffffff"
        colors.append(color)
    return colors


def build_background_from_colors(colors: list[str]) -> str:
    if not colors:
        return "#ffffff"

    span = len(colors)
    stops: list[str] = []
    for offset, color in enumerate(colors):
        start_pct = offset * 100 / span
        end_pct = (offset + 1) * 100 / span
        stops.append(f"{color} {start_pct:.6f}%")
        stops.append(f"{color} {end_pct:.6f}%")
    return f"linear-gradient(to right, {', '.join(stops)})"


def build_span_background(
    start_idx: int,
    span: int,
    all_dates: list[date],
    occupied_indices: set[int] | None = None,
    today_date: date | None = None,
) -> str:
    return build_background_from_colors(
        build_span_colors(start_idx, span, all_dates, occupied_indices, today_date)
    )


def render_span_attrs(
    start_idx: int,
    span: int,
    all_dates: list[date],
    occupied_indices: set[int] | None = None,
    today_date: date | None = None,
) -> str:
    colors = build_span_colors(start_idx, span, all_dates, occupied_indices, today_date)
    background = build_background_from_colors(colors)
    end_idx = start_idx + span - 1
    return (
        f'colspan="{span}" data-start-index="{start_idx}" data-end-index="{end_idx}" '
        f'data-original-colspan="{span}" data-bg-colors="{html.escape("|".join(colors))}" '
        f'style="--span-bg:{background};"'
    )


def build_weekend_overlay(all_dates: list[date]) -> str:
    if not all_dates:
        return "transparent"

    stops: list[str] = []
    for idx, day in enumerate(all_dates):
        color = HOLIDAY_FILL if is_holiday(day) else "transparent"
        start_pct = idx * 100 / len(all_dates)
        end_pct = (idx + 1) * 100 / len(all_dates)
        stops.append(f"{color} {start_pct:.6f}%")
        stops.append(f"{color} {end_pct:.6f}%")
    return f"linear-gradient(to right, {', '.join(stops)})"


def get_domain_class(domain: str) -> str:
    normalized = domain.strip().lower()
    if normalized == "crm":
        return "domain-crm"
    if normalized == "记账":
        return "domain-jizhang"
    if normalized == "税务":
        return "domain-shuiwu"
    return "domain-default"


def get_segment_class(segment_type: str) -> str:
    if segment_type == "case":
        return "segment-case"
    if segment_type == "execution":
        return "segment-execution"
    if segment_type == "expected":
        return "segment-expected"
    return "segment-default"


def render_task_card(person_task: PersonTask) -> str:
    task = person_task.task
    max_chars = max(8, min(42, person_task.span * 10))
    display_title = middle_ellipsis(task.title, max_chars)
    title = html.escape(display_title)
    task_id = html.escape(task.task_id or "-")
    source = html.escape(task.source or "-")
    version = html.escape(task.version or "-")
    remark = html.escape(task.remark or "").replace("\n", " ")
    parent_record = html.escape(task.parent_record or "-")
    tooltip = (
        f"任务编号：{task_id}\n"
        f"任务标题：{task.title}\n"
        f"阶段：{person_task.segment_label}\n"
        f"负责人：{person_task.owner}\n"
        f"任务来源：{task.source or '-'}\n"
        f"业务域：{task.domain or '-'}\n"
        f"测试类型：{task.test_type or '-'}\n"
        f"优先级：{task.priority or '-'}\n"
        f"所属版本：{task.version or '-'}\n"
        f"当前状态：{task.status or '-'}\n"
        f"开始：{person_task.start:%Y-%m-%d}\n"
        f"结束：{person_task.end:%Y-%m-%d}\n"
        f"父记录：{task.parent_record or '-'}\n"
        f"备注：{task.remark or '-'}"
    )
    segment_class = get_segment_class(person_task.segment_type)
    return f"""
    <div class="task-card" title="{html.escape(tooltip)}" data-task-id="{task_id}" data-title="{html.escape(task.title)}" data-source="{source}" data-version="{version}" data-parent="{parent_record}" data-remark="{remark}">
      <div class="task-segment {segment_class}">{html.escape(person_task.segment_label)}</div>
      <div class="task-main">{title}</div>
    </div>
    """


def render_matrix_rows(
    owner: str,
    lanes: list[list[PersonTask]],
    date_to_idx: dict[date, int],
    total_days: int,
    all_dates: list[date],
    today_date: date,
) -> str:
    if not lanes:
        empty_attrs = render_span_attrs(0, total_days, all_dates, set(), today_date)
        return f"""
        <tr class="person-row person-row-end" data-owner="{html.escape(owner)}" data-search="{html.escape(owner.lower())}">
          <th class="person-cell sticky-person">{html.escape(owner)}</th>
          <td class="day-cell empty-row-grid" {empty_attrs}>
            <div class="empty-row-label">没有已排期任务</div>
          </td>
        </tr>
        """

    rows = []
    search_blob = owner.lower()
    occupied_indices: set[int] = set()
    for lane in lanes:
        for item in lane:
            start_idx = date_to_idx[item.start]
            end_idx = date_to_idx[item.end]
            occupied_indices.update(range(start_idx, end_idx + 1))
            search_blob += " " + " ".join(
                [
                    item.task.task_id.lower(),
                    item.task.title.lower(),
                    item.task.source.lower(),
                    item.task.domain.lower(),
                    item.task.test_type.lower(),
                    item.task.status.lower(),
                    item.task.remark.lower(),
                ]
            )

    rowspan = len(lanes)
    for row_idx, lane in enumerate(lanes):
        current_col = 0
        cells = []
        for item in lane:
            start_idx = date_to_idx[item.start]
            if start_idx > current_col:
                empty_span = start_idx - current_col
                empty_attrs = render_span_attrs(
                    current_col,
                    empty_span,
                    all_dates,
                    occupied_indices,
                    today_date,
                )
                cells.append(
                    f'<td class="day-cell empty-cell" {empty_attrs}></td>'
                )
            task_attrs = render_span_attrs(
                start_idx,
                item.span,
                all_dates,
                occupied_indices,
                today_date,
            )
            cells.append(
                f'<td class="day-cell task-cell" {task_attrs}>{render_task_card(item)}</td>'
            )
            current_col = start_idx + item.span
        if current_col < total_days:
            tail_span = total_days - current_col
            tail_attrs = render_span_attrs(
                current_col,
                tail_span,
                all_dates,
                occupied_indices,
                today_date,
            )
            cells.append(
                f'<td class="day-cell empty-cell" {tail_attrs}></td>'
            )

        person_cell = ""
        if row_idx == 0:
            person_cell = (
                f'<th class="person-cell sticky-person" rowspan="{rowspan}">{html.escape(owner)}'
                f'<span class="lane-count">{rowspan} 行</span></th>'
            )

        rows.append(
            f"""
            <tr class="person-row{' person-row-end' if row_idx == rowspan - 1 else ''}" data-owner="{html.escape(owner)}" data-search="{html.escape(search_blob)}">
              {person_cell}
              {''.join(cells)}
            </tr>
            """
        )
    return "".join(rows)


def render_unscheduled_rows(unscheduled_by_person: dict[str, list[Task]]) -> str:
    rows = []
    for owner in sorted(unscheduled_by_person):
        for task in unscheduled_by_person[owner]:
            rows.append(
                f"""
                <tr>
                  <td>{html.escape(owner)}</td>
                  <td>{html.escape(task.task_id or "-")}</td>
                  <td>{html.escape(task.title)}</td>
                  <td>{html.escape(task.priority or "-")}</td>
                  <td>{html.escape(task.status or "-")}</td>
                  <td>未排期</td>
                  <td>{html.escape(task.remark or "-")}</td>
                </tr>
                """
            )
    return "".join(rows) if rows else '<tr><td colspan="7" class="empty-row">没有未排期任务</td></tr>'


def format_date_cell(value: date | None) -> str:
    return value.strftime("%Y-%m-%d") if value else ""


def task_unique_key(task: Task) -> tuple[object, ...]:
    return (
        task.task_id,
        task.title,
        tuple(task.owners),
        task.source,
        task.domain,
        task.test_type,
        task.priority,
        task.version,
        task.status,
        task.expected_test_date.isoformat() if task.expected_test_date else "",
        task.case_start.isoformat() if task.case_start else "",
        task.case_end.isoformat() if task.case_end else "",
        task.plan_start.isoformat() if task.plan_start else "",
        task.plan_end.isoformat() if task.plan_end else "",
        task.remark,
        task.parent_record,
    )


def dedupe_tasks(tasks: list[Task]) -> list[Task]:
    seen: set[tuple[object, ...]] = set()
    unique_tasks: list[Task] = []
    for task in tasks:
        key = task_unique_key(task)
        if key in seen:
            continue
        seen.add(key)
        unique_tasks.append(task)
    return unique_tasks


def get_window_scheduled_tasks(tasks: list[Task], window_start: date, window_end: date) -> list[Task]:
    scheduled_tasks: list[Task] = []
    for task in tasks:
        if is_pending_schedule(task):
            continue
        for start, end, _segment_type, _segment_label in build_task_segments(task):
            clipped_start = max(start, window_start)
            clipped_end = min(end, window_end)
            if clipped_start <= clipped_end:
                scheduled_tasks.append(task)
                break
    scheduled_tasks = dedupe_tasks(scheduled_tasks)
    scheduled_tasks.sort(key=lambda task: (task.task_id, task.title))
    return scheduled_tasks


def render_task_source_row(task: Task) -> str:
    cells = [
        task.task_id,
        task.title,
        "、".join(task.owners),
        task.domain,
        task.status,
        format_date_cell(task.expected_test_date),
        format_date_cell(task.case_start),
        format_date_cell(task.case_end),
        format_date_cell(task.plan_start),
        format_date_cell(task.plan_end),
        task.remark,
    ]
    return "<tr>" + "".join(f"<td>{html.escape(value or '-')}</td>" for value in cells) + "</tr>"


def render_pending_schedule_rows(tasks: list[Task]) -> str:
    pending_tasks = dedupe_tasks([task for task in tasks if is_pending_schedule(task)])
    pending_tasks.sort(
        key=lambda task: (
            "、".join(task.owners),
            priority_rank(task.priority),
            task.task_id,
            task.title,
        )
    )
    rows = [render_task_source_row(task) for task in pending_tasks]
    return "".join(rows) if rows else '<tr><td colspan="11" class="empty-row">没有待排期任务</td></tr>'


def render_scheduled_rows(tasks: list[Task], window_start: date, window_end: date) -> str:
    scheduled_tasks = get_window_scheduled_tasks(tasks, window_start, window_end)
    rows = [render_task_source_row(task) for task in scheduled_tasks]
    return "".join(rows) if rows else '<tr><td colspan="11" class="empty-row">没有已排期任务</td></tr>'


def render_template(context: dict[str, str]) -> str:
    template = TEMPLATE_PATH.read_text(encoding="utf-8")
    for key, value in context.items():
        template = template.replace("{{" + key + "}}", value)
    return template


def normalize_hex_color(value: str, fallback: str = "FFFFFF") -> str:
    raw = value.strip().lstrip("#")
    if len(raw) == 6 and all(char in "0123456789abcdefABCDEF" for char in raw):
        return f"FF{raw.upper()}"
    if len(raw) == 8 and all(char in "0123456789abcdefABCDEF" for char in raw):
        return raw.upper()
    fallback_raw = fallback.strip().lstrip("#")
    if len(fallback_raw) == 6 and all(char in "0123456789abcdefABCDEF" for char in fallback_raw):
        return f"FF{fallback_raw.upper()}"
    if len(fallback_raw) == 8 and all(char in "0123456789abcdefABCDEF" for char in fallback_raw):
        return fallback_raw.upper()
    return "FFFFFFFF"


def get_segment_fill(segment_type: str) -> PatternFill:
    if segment_type == "case":
        return PatternFill("solid", fgColor=normalize_hex_color("BACEFD"))
    if segment_type == "execution":
        return PatternFill("solid", fgColor=normalize_hex_color("3370FF"))
    if segment_type == "expected":
        return PatternFill("solid", fgColor=normalize_hex_color("3370FF"))
    return PatternFill("solid", fgColor=normalize_hex_color("3370FF"))


def get_segment_font(segment_type: str) -> Font:
    if segment_type == "case":
        return Font(name="Microsoft YaHei", size=9, color=normalize_hex_color("1F2329"), bold=True)
    return Font(name="Microsoft YaHei", size=9, color=normalize_hex_color("FFFFFF"), bold=True)


def get_current_month_range(today_date: date) -> tuple[date, date]:
    month_start = today_date.replace(day=1)
    return month_start, get_month_end(today_date)


def get_month_end(day: date) -> date:
    if day.month == 12:
        next_month_start = date(day.year + 1, 1, 1)
    else:
        next_month_start = date(day.year, day.month + 1, 1)
    return next_month_start - timedelta(days=1)


def get_latest_schedule_date(tasks: list[Task]) -> date | None:
    latest_date: date | None = None
    for task in tasks:
        for _start, end, _segment_type, _segment_label in build_task_segments(task):
            if latest_date is None or end > latest_date:
                latest_date = end
    return latest_date


def get_default_excel_range(tasks: list[Task], today_date: date) -> tuple[date, date]:
    excel_start, excel_end = get_current_month_range(today_date)
    latest_schedule_date = get_latest_schedule_date(tasks)
    if latest_schedule_date and latest_schedule_date > excel_end:
        excel_end = get_month_end(latest_schedule_date)
    return excel_start, excel_end


def get_default_page_range(today_date: date) -> tuple[date, date]:
    return today_date - timedelta(days=15), today_date + timedelta(days=30)


def resolve_date_window(
    default_start: date,
    default_end: date,
    start_value: str,
    end_value: str,
    start_arg_name: str = "--start-date",
    end_arg_name: str = "--end-date",
) -> tuple[date, date]:
    window_start = parse_date_arg(start_value, start_arg_name) or default_start
    window_end = parse_date_arg(end_value, end_arg_name) or default_end
    if window_end < window_start:
        window_start, window_end = window_end, window_start
    return window_start, window_end


def set_cell_border(cell, side: Side) -> None:
    cell.border = Border(top=side, bottom=side)


def get_duration_text(start: date | None, end: date | None) -> str:
    resolved_start, resolved_end = normalize_range(start, end)
    if not resolved_start or not resolved_end:
        return "-"
    duration = (resolved_end - resolved_start).days + 1
    return f"{duration}天"


def build_task_name_sheet(
    workbook: Workbook,
    tasks: list[Task],
    thin_grid: Side,
    header_fill: PatternFill,
    white_fill: PatternFill,
    header_font: Font,
    body_font: Font,
) -> None:
    sheet = workbook.create_sheet("任务名称维度")
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False

    headers = [
        "任务名称",
        "负责人",
        "任务状态",
        "测试用例耗时",
        "测试执行耗时",
        "用例开始时间",
        "用例结束时间",
        "测试开始时间",
        "测试结束时间",
    ]
    for col_idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        set_cell_border(cell, thin_grid)

    sorted_tasks = sorted(tasks, key=lambda task: (task.title, task.task_id))
    for row_idx, task in enumerate(sorted_tasks, start=2):
        values = [
            task.title,
            "、".join(task.owners) or "-",
            task.status or "-",
            get_duration_text(task.case_start, task.case_end),
            get_duration_text(task.plan_start, task.plan_end),
            format_date_cell(task.case_start) or "-",
            format_date_cell(task.case_end) or "-",
            format_date_cell(task.plan_start) or "-",
            format_date_cell(task.plan_end) or "-",
        ]
        sheet.row_dimensions[row_idx].height = 24
        for col_idx, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = white_fill
            cell.font = body_font
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            set_cell_border(cell, thin_grid)

    widths = [46, 18, 16, 14, 14, 16, 16, 16, 16]
    for col_idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col_idx)].width = width
    sheet.row_dimensions[1].height = 24
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(sorted_tasks) + 1, 1)}"


def append_task_detail_section(
    title: str,
    sheet,
    start_row: int,
    tasks: list[Task],
    thin_grid: Side,
    header_fill: PatternFill,
    white_fill: PatternFill,
    header_font: Font,
    body_font: Font,
) -> int:
    header_specs = [
        (1, 1, "任务编号"),
        (2, 4, "任务标题"),
        (5, 5, "负责人"),
        (6, 6, "业务域"),
        (7, 7, "当前状态"),
        (8, 8, "预计提测时间"),
        (9, 9, "用例开始时间"),
        (10, 10, "用例结束时间"),
        (11, 11, "计划开始时间"),
        (12, 12, "计划完成时间"),
        (13, 13, "备注"),
    ]

    title_end_col = 13
    sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=title_end_col)
    title_cell = sheet.cell(row=start_row, column=1, value=title)
    title_cell.fill = header_fill
    title_cell.font = Font(name="Microsoft YaHei", size=11, color=normalize_hex_color("334155"), bold=True)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    set_cell_border(title_cell, thin_grid)
    sheet.row_dimensions[start_row].height = 24

    header_row = start_row + 1
    for start_col, end_col, header in header_specs:
        if start_col != end_col:
            sheet.merge_cells(start_row=header_row, start_column=start_col, end_row=header_row, end_column=end_col)
        cell = sheet.cell(row=header_row, column=start_col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        set_cell_border(cell, thin_grid)
        for col_idx in range(start_col, end_col + 1):
            set_cell_border(sheet.cell(row=header_row, column=col_idx), thin_grid)

    sorted_tasks = sorted(
        dedupe_tasks(tasks),
        key=lambda task: (
            "、".join(task.owners),
            priority_rank(task.priority),
            task.task_id,
            task.title,
        ),
    )
    row_idx = header_row + 1
    for task in sorted_tasks:
        sheet.row_dimensions[row_idx].height = 24
        row_values = {
            1: task.task_id or "-",
            2: task.title or "-",
            5: "、".join(task.owners) or "-",
            6: task.domain or "-",
            7: task.status or "-",
            8: format_date_cell(task.expected_test_date) or "-",
            9: format_date_cell(task.case_start) or "-",
            10: format_date_cell(task.case_end) or "-",
            11: format_date_cell(task.plan_start) or "-",
            12: format_date_cell(task.plan_end) or "-",
            13: task.remark or "-",
        }
        sheet.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=4)
        for col_idx in range(1, title_end_col + 1):
            cell = sheet.cell(row=row_idx, column=col_idx)
            cell.fill = white_fill
            cell.font = body_font
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            set_cell_border(cell, thin_grid)
        for col_idx, value in row_values.items():
            sheet.cell(row=row_idx, column=col_idx, value=value)
        row_idx += 1

    if not sorted_tasks:
        sheet.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=title_end_col)
        empty_cell = sheet.cell(row=row_idx, column=1, value=f"没有{title}")
        empty_cell.fill = white_fill
        empty_cell.font = body_font
        empty_cell.alignment = Alignment(horizontal="center", vertical="center")
        set_cell_border(empty_cell, thin_grid)
        sheet.row_dimensions[row_idx].height = 24
        row_idx += 1

    return row_idx + 1


def export_excel_board(
    tasks: list[Task],
    output_path: Path,
    window_start: date,
    window_end: date,
    today_date: date,
) -> None:
    lanes_by_person, unscheduled_by_person, all_dates = build_schedule(tasks, window_start, window_end)
    people = sorted(lanes_by_person)
    date_to_idx = {day: idx for idx, day in enumerate(all_dates)}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人员任务看板"
    sheet.freeze_panes = "B3"
    sheet.sheet_view.showGridLines = False

    thin_grid = Side(style="thin", color="D8DEE9")
    person_border = Side(style="thin", color="000000")
    header_fill = PatternFill("solid", fgColor=normalize_hex_color("FFFFFF"))
    weekend_fill = PatternFill("solid", fgColor=normalize_hex_color(TASK_BAR_THEME["weekend"]))
    holiday_fill = PatternFill("solid", fgColor=normalize_hex_color(HOLIDAY_FILL))
    idle_fill = PatternFill("solid", fgColor=normalize_hex_color(TASK_BAR_THEME["idle"]))
    today_fill = PatternFill("solid", fgColor=normalize_hex_color(TASK_BAR_THEME["today"]))
    white_fill = PatternFill("solid", fgColor=normalize_hex_color("FFFFFF"))
    header_font = Font(name="Microsoft YaHei", size=10, color=normalize_hex_color("334155"), bold=True)
    body_font = Font(name="Microsoft YaHei", size=9, color=normalize_hex_color("1F2329"))

    sheet.cell(row=1, column=1, value="负责人")
    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    owner_header = sheet.cell(row=1, column=1)
    owner_header.fill = header_fill
    owner_header.font = header_font
    owner_header.alignment = Alignment(horizontal="center", vertical="center")
    set_cell_border(owner_header, thin_grid)

    month_start_col = 2
    current_month = all_dates[0].month if all_dates else None
    current_year = all_dates[0].year if all_dates else None
    for idx, day in enumerate(all_dates, start=2):
        if day.year != current_year or day.month != current_month:
            sheet.merge_cells(start_row=1, start_column=month_start_col, end_row=1, end_column=idx - 1)
            month_cell = sheet.cell(row=1, column=month_start_col)
            month_cell.value = f"{current_year}年{current_month}月"
            month_cell.fill = header_fill
            month_cell.font = header_font
            month_cell.alignment = Alignment(horizontal="center", vertical="center")
            current_year = day.year
            current_month = day.month
            month_start_col = idx

        date_cell = sheet.cell(row=2, column=idx, value=get_day_label(day))
        if is_holiday(day):
            date_cell.fill = holiday_fill
        elif day == today_date:
            date_cell.fill = today_fill
        elif day.weekday() >= 5:
            date_cell.fill = weekend_fill
        else:
            date_cell.fill = header_fill
        date_cell.font = header_font
        date_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        set_cell_border(date_cell, thin_grid)

    if all_dates:
        end_col = len(all_dates) + 1
        sheet.merge_cells(start_row=1, start_column=month_start_col, end_row=1, end_column=end_col)
        month_cell = sheet.cell(row=1, column=month_start_col)
        month_cell.value = f"{current_year}年{current_month}月"
        month_cell.fill = header_fill
        month_cell.font = header_font
        month_cell.alignment = Alignment(horizontal="center", vertical="center")
        for col_idx in range(2, end_col + 1):
            month_header = sheet.cell(row=1, column=col_idx)
            month_header.fill = header_fill
            month_header.font = header_font
            month_header.alignment = Alignment(horizontal="center", vertical="center")
            set_cell_border(month_header, thin_grid)

    row_idx = 3
    for owner in people:
        lanes = lanes_by_person.get(owner, [])
        if not lanes:
            lanes = [[]]
        owner_start_row = row_idx
        occupied_indices: set[int] = set()
        for lane in lanes:
            for item in lane:
                occupied_indices.update(range(date_to_idx[item.start], date_to_idx[item.end] + 1))

        for lane in lanes:
            sheet.row_dimensions[row_idx].height = 24
            for day_idx, day in enumerate(all_dates):
                col_idx = day_idx + 2
                cell = sheet.cell(row=row_idx, column=col_idx)
                if is_holiday(day):
                    cell.fill = holiday_fill
                elif day.weekday() >= 5:
                    cell.fill = weekend_fill
                elif day == today_date:
                    cell.fill = today_fill
                elif day_idx not in occupied_indices:
                    cell.fill = idle_fill
                else:
                    cell.fill = white_fill
                cell.font = body_font
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                set_cell_border(cell, thin_grid)

            for item in lane:
                start_col = date_to_idx[item.start] + 2
                end_col = date_to_idx[item.end] + 2
                if start_col != end_col:
                    sheet.merge_cells(start_row=row_idx, start_column=start_col, end_row=row_idx, end_column=end_col)
                task_cell = sheet.cell(row=row_idx, column=start_col)
                task_cell.value = f"{item.segment_label}  {item.task.title}"
                task_cell.fill = get_segment_fill(item.segment_type)
                task_cell.font = get_segment_font(item.segment_type)
                task_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                task_cell.comment = None
                for col_idx in range(start_col, end_col + 1):
                    set_cell_border(sheet.cell(row=row_idx, column=col_idx), thin_grid)

            row_idx += 1

        if owner_start_row != row_idx - 1:
            sheet.merge_cells(start_row=owner_start_row, start_column=1, end_row=row_idx - 1, end_column=1)
        owner_cell = sheet.cell(row=owner_start_row, column=1)
        owner_cell.value = owner
        owner_cell.font = Font(name="Microsoft YaHei", size=10, color=normalize_hex_color("1F2329"), bold=True)
        owner_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        owner_cell.fill = white_fill
        for border_row in range(owner_start_row, row_idx):
            for col_idx in range(1, len(all_dates) + 2):
                cell = sheet.cell(row=border_row, column=col_idx)
                current = cell.border
                cell.border = Border(
                    top=person_border if border_row == owner_start_row else current.top,
                    bottom=person_border if border_row == row_idx - 1 else current.bottom,
                )

    sheet.column_dimensions["A"].width = 14
    for col_idx in range(2, len(all_dates) + 2):
        sheet.column_dimensions[get_column_letter(col_idx)].width = 15
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 46
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(all_dates) + 1)}{max(row_idx - 1, 2)}"

    append_task_detail_section(
        "无负责人任务",
        sheet,
        row_idx + 2,
        [task for task in get_window_scheduled_tasks(tasks, window_start, window_end) if not task.owners],
        thin_grid,
        header_fill,
        white_fill,
        header_font,
        body_font,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def build_page(
    tasks: list[Task],
    input_path: Path,
    window_start: date | None = None,
    window_end: date | None = None,
) -> str:
    today_date = datetime.now().date()
    default_window_start, default_window_end = get_default_page_range(today_date)
    window_start = window_start or default_window_start
    window_end = window_end or default_window_end
    lanes_by_person, _unscheduled_by_person, all_dates = build_schedule(tasks, window_start, window_end)
    people = sorted(lanes_by_person)
    date_to_idx = {day: idx for idx, day in enumerate(all_dates)}
    total_lanes = sum(len(lanes_by_person.get(person, [])) or 1 for person in people)
    today = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    date_range_text = f"{window_start:%Y-%m-%d} ~ {window_end:%Y-%m-%d}"

    month_header_cells = render_month_headers(all_dates)
    header_cells = []
    for day in all_dates:
        classes = ["sticky-date"]
        if day.weekday() >= 5:
            classes.append("weekend")
        if day == today_date:
            classes.append("today-col")
        weekday = "一二三四五六日"[day.weekday()]
        header_cells.append(
            f'<th class="{" ".join(classes)}" data-date-index="{date_to_idx[day]}" data-date="{day.isoformat()}">'
            f'{format_day(day)}<span>周{weekday}</span></th>'
        )

    body_rows = []
    stats = []
    for person in people:
        lanes = lanes_by_person.get(person, [])
        task_count = sum(len(lane) for lane in lanes)
        stats.append(
            {
                "name": person,
                "lanes": len(lanes) or 1,
                "tasks": task_count,
            }
        )
        body_rows.append(render_matrix_rows(person, lanes, date_to_idx, len(all_dates), all_dates, today_date))

    stats_json = json.dumps(stats, ensure_ascii=False).replace("</", "<\\/")
    date_meta = [
        {
            "date": day.isoformat(),
            "month": f"{day.year}年{day.month}月",
            "day": format_day(day),
            "weekday": f"周{'一二三四五六日'[day.weekday()]}",
        }
        for day in all_dates
    ]
    date_meta_json = json.dumps(date_meta, ensure_ascii=False).replace("</", "<\\/")
    colgroup_html = render_colgroup(all_dates)
    weekend_overlay = build_weekend_overlay(all_dates)

    template_context = {
        "line": TASK_BAR_THEME["grid"],
        "line_strong": TASK_BAR_THEME["grid_strong"],
        "bar": TASK_BAR_THEME["bar"],
        "bar_deep": TASK_BAR_THEME["bar_deep"],
        "today_color": TASK_BAR_THEME["today"],
        "weekend": TASK_BAR_THEME["weekend"],
        "idle": TASK_BAR_THEME["idle"],
        "case_tag": TASK_BAR_THEME["case_tag"],
        "execution_tag": TASK_BAR_THEME["execution_tag"],
        "expected_tag": TASK_BAR_THEME["expected_tag"],
        "day_column_width": str(DAY_COLUMN_WIDTH),
        "person_column_width": str(PERSON_COLUMN_WIDTH),
        "weekend_overlay": weekend_overlay,
        "date_count": str(len(all_dates)),
        "input_name": html.escape(input_path.name),
        "people_count": str(len(people)),
        "date_range": html.escape(date_range_text),
        "default_filter_start": window_start.isoformat(),
        "default_filter_end": window_end.isoformat(),
        "total_lanes": str(total_lanes),
        "generated_at": html.escape(today),
        "colgroup_html": colgroup_html,
        "month_header_cells": month_header_cells,
        "header_cells": "".join(header_cells),
        "body_rows": "".join(body_rows),
        "pending_schedule_rows": render_pending_schedule_rows(tasks),
        "scheduled_rows": render_scheduled_rows(tasks, window_start, window_end),
        "stats_json": stats_json,
        "date_meta_json": date_meta_json,
    }
    return render_template(template_context)


def main() -> None:
    args = parse_args()
    input_path = args.input.resolve()
    output_path = args.output.resolve()
    excel_output_path = args.excel_output.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    tasks = load_tasks(input_path)
    today_date = datetime.now().date()
    custom_start_value = args.start_date or args.excel_start
    custom_end_value = args.end_date or args.excel_end
    has_custom_range = bool(custom_start_value or custom_end_value)
    page_start, page_end = resolve_date_window(
        *get_default_page_range(today_date),
        custom_start_value,
        custom_end_value,
    )

    if has_custom_range:
        page = build_page(tasks, input_path, page_start, page_end)
    else:
        page = build_page(tasks, input_path)
    output_path.write_text(page, encoding="utf-8")
    index_output_path = output_path.with_name(INDEX_OUTPUT_NAME)
    if output_path != index_output_path:
        shutil.copyfile(output_path, index_output_path)

    if has_custom_range:
        excel_start, excel_end = page_start, page_end
    else:
        excel_start, excel_end = get_default_excel_range(tasks, today_date)
    export_excel_board(tasks, excel_output_path, excel_start, excel_end, today_date)

    print(f"输入文件: {input_path}")
    print(f"输出 HTML: {output_path}")
    if output_path != index_output_path:
        print(f"同步首页: {index_output_path}")
    print(f"输出 Excel: {excel_output_path}")
    print(f"HTML 日期范围: {page_start:%Y-%m-%d} ~ {page_end:%Y-%m-%d}")
    print(f"Excel 日期范围: {excel_start:%Y-%m-%d} ~ {excel_end:%Y-%m-%d}")
    print(f"任务总数: {len(tasks)}")


if __name__ == "__main__":
    main()

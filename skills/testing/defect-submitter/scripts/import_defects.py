# -*- coding: utf-8 -*-
"""从 Excel 批量导入缺陷到缺陷管理平台。

用法：
    python import_defects.py                # 正式导入（读 缺陷导入模板.xlsx）
    python import_defects.py --dry-run      # 只校验数据，不调接口
    python import_defects.py --file 其他文件.xlsx

幂等：「导入状态」列以"成功"开头的行会跳过，可重复执行。

平台差异全部在 defect_client.py，本脚本不直接调任何平台 API。
"""

import argparse
import sys
import time
from pathlib import Path

import openpyxl

from defect_client import DefectApiError, load_config, client_from_config

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_EXCEL = BASE_DIR / "缺陷导入模板.xlsx"

# Excel 列定义（A~I）
COL_NO, COL_SUBJECT, COL_DESC, COL_SEVERITY, COL_PRIORITY, COL_ASSIGNEE, COL_SPRINT, COL_VERSION, COL_STATUS = range(1, 10)


def map_option(cfg, field_name, raw_value):
    """按 config.field_map 把 Excel 里的中文选项映射为 (fieldId, optionId)。
    返回 None 表示该字段未配置或值为空（跳过）；映射不到时返回错误信息。"""
    field_cfg = cfg.get("field_map", {}).get(field_name)
    if raw_value is None or str(raw_value).strip() == "":
        return None, None
    value = str(raw_value).strip()
    if not field_cfg or not field_cfg.get("field_id"):
        return None, f"[{field_name}] 未在 config.field_map 中配置字段，值「{value}」被忽略"
    option_id = field_cfg.get("options", {}).get(value)
    if not option_id:
        return None, f"[{field_name}] 值「{value}」在 config 中无映射，被忽略"
    return (field_cfg["field_id"], option_id), None


def build_row_payload(cfg, sheet_name, row):
    """把一行 Excel 数据组装为 create_issue 参数。返回 (kwargs, warnings, error)。"""
    warnings = []
    subject = row[COL_SUBJECT]
    if subject is None or str(subject).strip() == "":
        return None, [], "缺少缺陷标题"
    subject = str(subject).strip()

    # 指派人：姓名 -> 平台用户标识，兜底默认指派人
    assignee_name = str(row[COL_ASSIGNEE]).strip() if row[COL_ASSIGNEE] else ""
    assigned_to = cfg.get("users", {}).get(assignee_name)
    if not assigned_to:
        if assignee_name:
            warnings.append(f"指派人「{assignee_name}」无映射，使用默认指派人")
        assigned_to = cfg.get("default_assignee")
    if not assigned_to:
        return None, warnings, "指派人为空且未配置 default_assignee"

    custom = {}
    # 严重程度 / 优先级
    for field_name, col in (("严重程度", COL_SEVERITY), ("优先级", COL_PRIORITY)):
        mapped, warn = map_option(cfg, field_name, row[col])
        if warn:
            warnings.append(warn)
        if mapped:
            custom[mapped[0]] = mapped[1]
    # 模块 = sheet 名
    mapped, warn = map_option(cfg, "模块", sheet_name)
    if warn and cfg.get("field_map", {}).get("模块", {}).get("field_id"):
        warnings.append(warn)
    if mapped:
        custom[mapped[0]] = mapped[1]

    kwargs = {
        "title": subject,
        "assignee": assigned_to,
        "description": str(row[COL_DESC]).strip() if row[COL_DESC] else None,
        "custom_fields": custom or None,
    }

    # 迭代 / 版本（名称 -> id，可空；平台不支持时适配层忽略）
    sprint_name = str(row[COL_SPRINT]).strip() if row[COL_SPRINT] else ""
    if sprint_name:
        sprint_id = cfg.get("sprints", {}).get(sprint_name)
        if sprint_id:
            kwargs["sprint"] = sprint_id
        else:
            warnings.append(f"迭代「{sprint_name}」无映射，已忽略")
    version_name = str(row[COL_VERSION]).strip() if row[COL_VERSION] else ""
    if version_name:
        version_id = cfg.get("versions", {}).get(version_name)
        if version_id:
            kwargs["versions"] = [version_id]
        else:
            warnings.append(f"版本「{version_name}」无映射，已忽略")

    return kwargs, warnings, None


def main():
    parser = argparse.ArgumentParser(description="从 Excel 批量导入缺陷")
    parser.add_argument("--file", default=str(DEFAULT_EXCEL), help="Excel 文件路径")
    parser.add_argument("--dry-run", action="store_true", help="只校验数据，不调用接口")
    args = parser.parse_args()

    cfg = load_config()
    excel_path = Path(args.file)
    if not excel_path.exists():
        print(f"文件不存在: {excel_path}")
        sys.exit(1)

    client = None
    if not args.dry_run:
        for key in ("project", "issue_type"):
            if not cfg.get(key):
                print(f"config.json 缺少配置项: {key}")
                sys.exit(1)
        try:
            client = client_from_config(cfg)
        except DefectApiError as e:
            print(e)
            sys.exit(1)

    wb = openpyxl.load_workbook(excel_path)
    interval = float(cfg.get("request_interval", 0.3))
    summary = {}
    dirty = False

    for ws in wb.worksheets:
        ok = fail = skip = 0
        for r in range(2, ws.max_row + 1):
            row = {c: ws.cell(row=r, column=c).value for c in range(1, 10)}
            if row[COL_NO] is None and row[COL_SUBJECT] is None:
                continue  # 空行
            status = str(row[COL_STATUS] or "")
            if status.startswith("成功"):
                skip += 1
                continue

            kwargs, warnings, error = build_row_payload(cfg, ws.title, row)
            for w in warnings:
                print(f"  [提醒] {ws.title} 第{r}行: {w}")
            if error:
                ws.cell(row=r, column=COL_STATUS, value=f"失败: {error}")
                dirty = True
                fail += 1
                continue
            if args.dry_run:
                print(f"  [预检] {ws.title} 第{r}行 OK: {kwargs['title'][:30]}")
                ok += 1
                continue

            try:
                issue_id = client.create_issue(
                    project=cfg["project"],
                    type_id=cfg["issue_type"],
                    **kwargs,
                )
                ws.cell(row=r, column=COL_STATUS, value=f"成功 {issue_id}")
                ok += 1
                print(f"  [成功] {ws.title} 第{r}行 -> {issue_id} {kwargs['title'][:30]}")
            except DefectApiError as e:
                ws.cell(row=r, column=COL_STATUS, value=f"失败: {e}")
                fail += 1
                print(f"  [失败] {ws.title} 第{r}行: {e}")
            dirty = True
            time.sleep(interval)

        if ok or fail or skip:
            summary[ws.title] = (ok, fail, skip)

    if dirty and not args.dry_run:
        wb.save(excel_path)
        print(f"\n结果已回写: {excel_path}")

    print("\n===== 导入汇总 =====")
    print(f"{'Sheet':<14}{'成功':>6}{'失败':>6}{'跳过':>6}")
    total = [0, 0, 0]
    for name, (ok, fail, skip) in summary.items():
        print(f"{name:<14}{ok:>6}{fail:>6}{skip:>6}")
        total[0] += ok; total[1] += fail; total[2] += skip
    print(f"{'合计':<14}{total[0]:>6}{total[1]:>6}{total[2]:>6}")


if __name__ == "__main__":
    main()

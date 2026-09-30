#!/usr/bin/env python3
"""将 BOSS直聘 抓取结果（岗位列表 + 详情）导出为 Excel，岗位和详情按 job_id 对应。"""

import argparse
import json
import os
import re
import sys
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("请先安装 openpyxl: pip install openpyxl")
    sys.exit(1)


DATA_DIR = os.path.expanduser("~/.boss-zhipin-scraper/job-result")


def find_latest_files(data_dir):
    """找到最新的 jobs 和 details JSON 文件。"""
    jobs_files = []
    details_files = []
    if not os.path.isdir(data_dir):
        print(f"目录不存在: {data_dir}")
        return None, None

    for f in os.listdir(data_dir):
        if f.startswith("boss_jobs_") and f.endswith(".json"):
            jobs_files.append(f)
        elif f.startswith("boss_details_") and f.endswith(".json"):
            details_files.append(f)

    if not jobs_files:
        print(f"未找到 boss_jobs_*.json 文件: {data_dir}")
        return None, None

    jobs_files.sort(reverse=True)
    details_files.sort(reverse=True)

    jobs_path = os.path.join(data_dir, jobs_files[0])
    ts = jobs_files[0].replace("boss_jobs_", "").replace(".json", "")
    matched = [f for f in details_files if ts in f]
    if matched:
        details_path = os.path.join(data_dir, matched[0])
    elif details_files:
        details_path = os.path.join(data_dir, details_files[0])
    else:
        details_path = None

    return jobs_path, details_path


def clean_jd(jd_text):
    """清理 JD 文本。"""
    if not jd_text:
        return ""
    text = jd_text
    text = re.sub(r"微\s*信\s*扫码分享\s*举报\s*", "", text)
    text = re.split(r"BOSS\s*安全提示", text)[0]
    text = re.split(r"公司介绍", text)[0]
    text = re.split(r"工商信息", text)[0]
    text = re.split(r"工作地址", text)[0]
    text = re.split(r"更多职位", text)[0]
    text = re.split(r"看过该职位的人还看了", text)[0]
    text = re.split(r"精选职位", text)[0]
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    return text.strip()


def extract_company_name(job):
    return job.get("brand_name", job.get("company_name", ""))


def merge_jobs_with_details(jobs_data, details_data):
    """按 job_id 合并岗位列表和详情。"""
    detail_map = {}
    for d in details_data:
        if isinstance(d, dict):
            jid = d.get("job_id") or d.get("encrypt_job_id") or ""
            detail_map[jid] = d
            link = d.get("job_link") or d.get("link") or ""
            if link:
                detail_map[link] = d

    merged = []
    jobs_list = jobs_data.get("jobs", []) if isinstance(jobs_data, dict) else jobs_data

    for job in jobs_list:
        jid = job.get("job_id") or job.get("encrypt_job_id") or ""
        link = job.get("job_link") or ""
        detail = detail_map.get(jid) or detail_map.get(link, {})

        merged.append({
            "序号": len(merged) + 1,
            "职位名称": job.get("title", ""),
            "薪资": job.get("salary", ""),
            "地点": job.get("location", ""),
            "公司名称": detail.get("company", extract_company_name(job)),
            "公司规模": job.get("company_scale", ""),
            "公司阶段": job.get("company_stage", ""),
            "行业": job.get("company_industry", ""),
            "标签": job.get("tags", job.get("job_labels", "")),
            "技能": job.get("skills", ""),
            "福利": job.get("welfare", ""),
            "BOSS": job.get("boss_name", ""),
            "BOSS职位": job.get("boss_title", ""),
            "学历经验": detail.get("tags_list", "") or job.get("tags", ""),
            "JD详情": clean_jd(detail.get("jd", "")),
            "JD链接": job.get("job_link", ""),
        })

    return merged


def write_excel(merged, out_path):
    """将合并数据写入格式化的 Excel 文件。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "岗位详情"

    headers = [
        "序号", "职位名称", "薪资", "地点", "公司名称", "公司规模",
        "公司阶段", "行业", "标签", "技能", "福利",
        "BOSS", "BOSS职位", "学历经验", "JD详情", "JD链接"
    ]

    # 样式定义
    h_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    h_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    h_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c_font = Font(name="微软雅黑", size=10)
    c_align = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    alt_fill = PatternFill(start_color="F2F7FB", end_color="F2F7FB", fill_type="solid")

    # 写表头
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = h_font
        cell.fill = h_fill
        cell.alignment = h_align
        cell.border = thin_border

    # 写数据
    for row_idx, record in enumerate(merged, 2):
        for col_idx, h in enumerate(headers, 1):
            val = record.get(h, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = c_font
            cell.alignment = c_align
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    # 列宽
    col_widths = {
        "序号": 6, "职位名称": 28, "薪资": 12, "地点": 14,
        "公司名称": 24, "公司规模": 14, "公司阶段": 10, "行业": 16,
        "标签": 24, "技能": 30, "福利": 24,
        "BOSS": 12, "BOSS职位": 14, "学历经验": 14,
        "JD详情": 80, "JD链接": 40,
    }
    for col_idx, h in enumerate(headers, 1):
        w = col_widths.get(h, 18)
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    # 冻结首行 + 自动筛选
    ws.freeze_panes = "A2"
    if merged:
        ws.auto_filter.ref = ws.dimensions

    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(description="将 BOSS直聘 抓取结果导出为 Excel")
    parser.add_argument("--input-jobs", help="jobs JSON 文件路径")
    parser.add_argument("--input-details", help="details JSON 文件路径")
    parser.add_argument("--output", help="输出 Excel 文件路径")
    parser.add_argument("--data-dir", default=DATA_DIR, help=f"数据目录（默认 {DATA_DIR}）")
    args = parser.parse_args()

    # 确定文件路径
    if args.input_jobs:
        jobs_path = args.input_jobs
        details_path = args.input_details or ""
    else:
        jp, dp = find_latest_files(args.data_dir)
        if not jp:
            sys.exit(1)
        jobs_path = jp
        details_path = dp or ""

    print(f"读取岗位列表: {jobs_path}")
    with open(jobs_path, encoding="utf-8") as f:
        jobs_data = json.load(f)

    if details_path and os.path.isfile(details_path):
        print(f"读取详情: {details_path}")
        with open(details_path, encoding="utf-8") as f:
            details_data = json.load(f)
        print(f"  详情记录数: {len(details_data)}")
    else:
        details_data = []
        print("未找到详情文件，仅导出岗位列表")

    # 合并
    merged = merge_jobs_with_details(jobs_data, details_data)

    if not merged:
        print("没有数据可导出")
        sys.exit(0)

    # 输出路径
    if args.output:
        out_path = args.output
    else:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_dir = DATA_DIR
        os.makedirs(out_dir, exist_ok=True)
        out_path = os.path.join(out_dir, f"boss_jobs_merged_{ts}.xlsx")

    print(f"\n合并记录数: {len(merged)}")
    print(f"导出 Excel: {out_path}")

    write_excel(merged, out_path)
    print("✅ 导出完成!")


if __name__ == "__main__":
    main()

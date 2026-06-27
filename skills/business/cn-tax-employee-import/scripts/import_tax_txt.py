#!/usr/bin/env python
"""Import Chinese tax export .txt data into employee/payroll Excel templates."""

from __future__ import annotations

import argparse
import glob
import json
import re
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any


CERT_TYPE = {201: "居民身份证", "201": "居民身份证"}
NATION = {156: "中国", "156": "中国"}
SEX = {1: "男", "1": "男", 2: "女", "2": "女"}
YESNO = {1: "是", "1": "是", True: "是", 0: "否", "0": "否", False: "否", "": ""}
EMPLOYMENT_TYPE = {10: "雇员", "10": "雇员"}
JOB = {0: "普通", "0": "普通", 1: "高层", "1": "高层"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--txt", nargs="+", required=True, help="Source .txt files or glob patterns")
    parser.add_argument("--personnel-template", help="人员信息采集导入模板.xls")
    parser.add_argument("--salary-template", help="工资表模板.xlsx")
    parser.add_argument("--tax-db", help="个税数据库.xlsx")
    parser.add_argument("--no-backup", action="store_true", help="Do not create timestamped backups")
    return parser.parse_args()


def expand_txt(patterns: list[str]) -> list[Path]:
    paths: list[Path] = []
    for pattern in patterns:
        matches = glob.glob(pattern)
        if matches:
            paths.extend(Path(m) for m in matches)
        else:
            paths.append(Path(pattern))
    unique: list[Path] = []
    seen = set()
    for path in paths:
        resolved = path.resolve()
        if resolved not in seen:
            seen.add(resolved)
            unique.append(path)
    return unique


def parse_tax_txt(path: Path) -> tuple[str, list[dict[str, Any]], list[dict[str, Any]]]:
    text = path.read_text(encoding="utf-8-sig")
    match = re.search(r"员工信息：\s*(\[.*?\])\s*工资信息：\s*(\[.*\])\s*$", text, re.S)
    if not match:
        raise ValueError(f"无法识别文件结构: {path}")
    company = path.stem
    employees = json.loads(match.group(1))
    salaries = json.loads(match.group(2))
    return company, employees, salaries


def make_backup(paths: list[Path]) -> Path | None:
    existing = [p for p in paths if p]
    if not existing:
        return None
    backup_dir = Path.cwd() / ("备份_个税导入_" + datetime.now().strftime("%Y%m%d_%H%M%S"))
    backup_dir.mkdir(exist_ok=True)
    for path in existing:
        shutil.copy2(path, backup_dir / path.name)
    return backup_dir


def as_text(value: Any) -> str:
    return "" if value is None else str(value)


def personnel_value(header: str, emp: dict[str, Any]) -> Any:
    h = (header or "").replace("*", "").strip()
    values = {
        "工号": emp.get("GH", ""),
        "姓名": emp.get("XM", ""),
        "证件类型": CERT_TYPE.get(emp.get("ZZLX"), emp.get("ZZLX", "")),
        "证件号码": emp.get("ZZHM", ""),
        "国籍(地区)": NATION.get(emp.get("GJ"), emp.get("GJ", "")),
        "性别": SEX.get(emp.get("XB"), emp.get("XB", "")),
        "出生日期": emp.get("CSNY", ""),
        "是否高级专家": "否" if emp.get("ZW", 0) in (0, "0", "", None) else "",
        "任职受雇从业类型": EMPLOYMENT_TYPE.get(emp.get("RZSGLX"), emp.get("RZSGLX", "")),
        "其他情况说明": emp.get("QTQKSM", ""),
        "入职年度就业情形": emp.get("RZNDJYQK", ""),
        "手机号码": emp.get("LXDH", ""),
        "任职受雇从业日期": emp.get("RZSGRQ", ""),
        "离职日期": emp.get("LZRQ", ""),
        "是否离职后补发工资": YESNO.get(emp.get("SFBFGZ"), emp.get("SFBFGZ", "")),
        "补发税款所属月份": emp.get("BFYF", ""),
        "是否残疾": YESNO.get(emp.get("SFCJ"), emp.get("SFCJ", "")),
        "是否烈属": YESNO.get(emp.get("SFLS"), emp.get("SFLS", "")),
        "是否孤老": YESNO.get(emp.get("SFGL"), emp.get("SFGL", "")),
        "残疾证件类型": emp.get("CJZJLX", ""),
        "残疾证号": emp.get("CJZH", ""),
        "烈属证号": emp.get("LSZH", ""),
        "是否扣除减除费用": YESNO.get(emp.get("SFKCJCFY"), emp.get("SFKCJCFY", "")),
        "备注": emp.get("BZ", ""),
        "电子邮箱": emp.get("DZYX", ""),
        "学历": emp.get("XL", ""),
        "开户银行": emp.get("KHYH", ""),
        "银行账号": emp.get("YHZH", ""),
        "开户行省份": emp.get("KHYH_SHENG", ""),
        "职务": JOB.get(emp.get("ZW"), ""),
        "户籍所在地（省）": emp.get("HJSZD_SHENG", ""),
        "户籍所在地（市）": emp.get("HJSZD_SHI", ""),
        "户籍所在地（区县）": emp.get("HJSZD_QX", ""),
        "户籍所在地（详细地址）": emp.get("HJSZD_XXDZ", ""),
        "经常居住地（省）": emp.get("JCJZD_SHENG", ""),
        "经常居住地（市）": emp.get("JCJZD_SHI", ""),
        "经常居住地（区县）": emp.get("JCJZD_QX", ""),
        "经常居住地（详细地址）": emp.get("JCJZD_XXDZ", ""),
        "联系地址（省）": emp.get("LXDZ_SHENG", ""),
        "联系地址（市）": emp.get("LXDZ_SHI", ""),
        "联系地址（区县）": emp.get("LXDZ_QX", ""),
        "联系地址（详细地址）": emp.get("LXDZ_XXDZ", ""),
    }
    if h == "个人投资额":
        return emp.get("GRTZE", "") if emp.get("GRTZE") not in (0, "0", "", None) else ""
    if h == "个人投资比例(%)":
        return emp.get("GRTZBL", "") if emp.get("GRTZBL") not in (0, "0", "", None) else ""
    return values.get(h, "")


def write_personnel_xls(path: Path, records: list[tuple[str, list[dict[str, Any]], list[dict[str, Any]]]]) -> int:
    try:
        import win32com.client as win32
    except Exception as exc:
        raise RuntimeError("写入 .xls 需要 Windows Excel COM（pywin32 + Excel）。") from exc

    employees = [emp for _company, emps, _salaries in records for emp in emps]
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(str(path.resolve()), ReadOnly=False, UpdateLinks=False)
        ws = wb.Worksheets("人员信息")
        used_cols = ws.UsedRange.Columns.Count
        headers = [ws.Cells(1, c).Value for c in range(1, used_cols + 1)]
        ws.Range(ws.Cells(2, 1), ws.Cells(1000, used_cols)).ClearContents()

        text_columns = {1, 4, 12, 20, 21, 22, 33, 46, 49}
        date_columns = {7, 13, 14, 30, 31}
        for col in text_columns:
            ws.Range(ws.Cells(2, col), ws.Cells(1000, col)).NumberFormat = "@"
        for col in date_columns:
            ws.Range(ws.Cells(2, col), ws.Cells(1000, col)).NumberFormat = "yyyy-mm-dd"

        for row, emp in enumerate(employees, start=2):
            for col, header in enumerate(headers, start=1):
                value = personnel_value(header, emp)
                if col in text_columns and value != "":
                    value = as_text(value)
                ws.Cells(row, col).Value = value
        wb.Save()
        wb.Close(False)
    finally:
        excel.Quit()
    return len(employees)


def salary_value(header: str, employee: dict[str, Any], salary: dict[str, Any]) -> Any:
    h = (header or "").replace("*", "").strip()
    values = {
        "工号": employee.get("GH", ""),
        "姓名": employee.get("XM", ""),
        "证件类型": CERT_TYPE.get(employee.get("ZZLX"), employee.get("ZZLX", "")),
        "证件号码": employee.get("ZZHM", ""),
        "本期收入": salary.get("SRE", 0),
        "本期免税收入": salary.get("MSSD", 0),
        "应补（退）税额": salary.get("YKJSE", 0),
        "基本养老保险费": salary.get("JBYLAOBXF", 0),
        "基本医疗保险费": salary.get("JBYLBXF", 0),
        "失业保险费": salary.get("SYBXF", 0),
        "住房公积金": salary.get("ZFGJJ", 0),
        "累计收入额": salary.get("SRE", 0),
        "累计免税收入": salary.get("MSSD", 0),
        "累计减除费用": salary.get("KCJJCHJ", 0),
        "累计专项扣除": salary.get("BQZXKCHJ", 0),
        "累计子女教育": salary.get("ZNJYZCLJ", 0),
        "累计继续教育": salary.get("JXJYZCLJ", 0),
        "累计住房贷款利息": salary.get("ZFDKLXZCLJ", 0),
        "累计住房租金": salary.get("ZFZJZCLJ", 0),
        "累计赡养老人": salary.get("SYLRZCLJ", 0),
        "累计3岁以下婴幼儿照护": salary.get("YYZHZCLJ", 0),
        "累计个人养老金": salary.get("GRYLJ", 0),
        "企业(职业)年金": salary.get("NJ", 0),
        "商业健康保险": salary.get("SYJKBX", 0),
        "税延养老保险": salary.get("SYYLBX", 0),
        "交通费": salary.get("JTF", 0),
        "通讯费": salary.get("TXF", 0),
        "律师办案费用": salary.get("BAF", 0),
        "西藏附加减除费用": salary.get("XZFJJCFY", 0),
        "准予扣除的捐赠额": salary.get("ZYKCDJZE", 0),
        "减免税额": salary.get("JMSE", 0),
    }
    return values.get(h, "")


def write_salary_xlsx(path: Path, records: list[tuple[str, list[dict[str, Any]], list[dict[str, Any]]]]) -> int:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]

    row = 2
    count = 0
    for _company, employees, salaries in records:
        emp_by_nsrid = {emp.get("NSRID"): emp for emp in employees}
        for salary in salaries:
            employee = emp_by_nsrid.get(salary.get("NSRID"), {})
            for col, header in enumerate(headers, start=1):
                ws.cell(row, col, salary_value(header, employee, salary))
            row += 1
            count += 1
    wb.save(path)
    return count


def remove_generated_columns(ws: Any) -> None:
    for col in range(ws.max_column, 0, -1):
        if ws.cell(1, col).value == "导入数据":
            ws.delete_cols(col, 1)


def add_dict_columns(ws: Any, columns: list[tuple[str, str, dict[str, Any]]]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    remove_generated_columns(ws)
    start = ws.max_column + 1
    header_fill = PatternFill("solid", fgColor="D9EAD3")
    sub_fill = PatternFill("solid", fgColor="EAF4E5")
    for offset, (company, label, data) in enumerate(columns):
        col = start + offset
        ws.cell(1, col, "导入数据")
        ws.cell(2, col, f"{company}-{label}")
        ws.cell(1, col).fill = header_fill
        ws.cell(2, col).fill = sub_fill
        ws.cell(1, col).font = Font(bold=True)
        ws.cell(2, col).font = Font(bold=True)
        ws.cell(1, col).alignment = Alignment(horizontal="center")
        ws.cell(2, col).alignment = Alignment(horizontal="center")
        for row in range(3, ws.max_row + 1):
            field = ws.cell(row, 1).value
            if field in data:
                ws.cell(row, col, data[field])


def write_tax_db_xlsx(path: Path, records: list[tuple[str, list[dict[str, Any]], list[dict[str, Any]]]]) -> tuple[int, int]:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    employee_cols: list[tuple[str, str, dict[str, Any]]] = []
    salary_cols: list[tuple[str, str, dict[str, Any]]] = []
    company_cols: list[tuple[str, str, dict[str, Any]]] = []
    for company, employees, salaries in records:
        for idx, emp in enumerate(employees, 1):
            employee_cols.append((company, f"员工{idx}-{emp.get('XM', '')}", emp))
        for idx, salary in enumerate(salaries, 1):
            label = as_text(salary.get("NSRID", ""))[:8] or f"工资{idx}"
            salary_cols.append((company, f"工资{idx}-{label}", salary))
        kjywrids = sorted(
            {emp.get("KJYWRID") for emp in employees if emp.get("KJYWRID") is not None}
            | {salary.get("KJYWRID") for salary in salaries if salary.get("KJYWRID") is not None}
        )
        company_cols.append((company, "扣缴义务人", {"KJYWRMC": company, "KJYWRID": kjywrids[0] if kjywrids else ""}))

    add_dict_columns(wb["人员表"], employee_cols)
    add_dict_columns(wb["工资薪金表"], salary_cols)
    if "扣缴义务人信息表" in wb.sheetnames:
        add_dict_columns(wb["扣缴义务人信息表"], company_cols)
    wb.save(path)
    return len(employee_cols), len(salary_cols)


def main() -> None:
    args = parse_args()
    outputs = [Path(p) for p in [args.personnel_template, args.salary_template, args.tax_db] if p]
    if not outputs:
        raise SystemExit("至少指定一个输出：--personnel-template、--salary-template 或 --tax-db")

    txt_paths = expand_txt(args.txt)
    records = [parse_tax_txt(path) for path in txt_paths]

    if not args.no_backup:
        backup_dir = make_backup(outputs)
        if backup_dir:
            print(f"备份目录: {backup_dir}")

    if args.personnel_template:
        count = write_personnel_xls(Path(args.personnel_template), records)
        print(f"人员信息采集导入模板写入: {count} 人")
    if args.salary_template:
        count = write_salary_xlsx(Path(args.salary_template), records)
        print(f"工资表模板写入: {count} 条")
    if args.tax_db:
        emp_count, salary_count = write_tax_db_xlsx(Path(args.tax_db), records)
        print(f"个税数据库写入: 人员 {emp_count} 条, 工资 {salary_count} 条")


if __name__ == "__main__":
    main()
